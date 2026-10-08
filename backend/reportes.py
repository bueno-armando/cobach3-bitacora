import io
import os
import copy
import datetime
from typing import Optional, List
from sqlalchemy.orm import Session, joinedload
from fastapi import HTTPException
import openpyxl

from backend.models import Activo, Resguardante

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_PATH = os.path.join(BASE_DIR, "RESGUARDO DE BIENES MUEBLES (FORMATO).xlsx")


def generar_resguardo_oficial_excel(
    db: Session,
    resguardante_id: Optional[int] = None,
    activo_ids: Optional[List[int]] = None
) -> io.BytesIO:
    """
    Genera el formato oficial institucional FOR-DAD_06 (Resguardo de Bienes Muebles)
    utilizando la plantilla Excel oficial de COBACH.
    """
    if not os.path.exists(TEMPLATE_PATH):
        raise HTTPException(status_code=500, detail="No se encontró la plantilla de resguardo oficial en el servidor.")

    # 1. Obtener los activos correspondientes
    query = db.query(Activo).options(
        joinedload(Activo.resguardante),
        joinedload(Activo.ubicacion)
    ).filter(Activo.deleted_at == None)

    nombre_resguardante = "RESGUARDO GENERAL"
    puesto_resguardante = ""

    if resguardante_id:
        resg = db.query(Resguardante).filter(Resguardante.id == resguardante_id).first()
        if not resg:
            raise HTTPException(status_code=404, detail="Resguardante no encontrado")
        nombre_resguardante = resg.nombre.strip().upper()
        puesto_resguardante = (resg.puesto or "").strip().upper()
        activos = query.filter(Activo.resguardante_id == resguardante_id).order_by(Activo.id.asc()).all()
    elif activo_ids:
        activos = query.filter(Activo.id.in_(activo_ids)).order_by(Activo.id.asc()).all()
        # Si todos comparten el mismo resguardante, usarlo
        resg_set = {a.resguardante for a in activos if a.resguardante}
        if len(resg_set) == 1:
            r = list(resg_set)[0]
            nombre_resguardante = r.nombre.strip().upper()
            puesto_resguardante = (r.puesto or "").strip().upper()
    else:
        raise HTTPException(status_code=400, detail="Debe especificar un resguardante_id o una lista de activo_ids.")

    if not activos:
        raise HTTPException(status_code=400, detail="No se encontraron activos para generar el resguardo.")

    # 2. Cargar plantilla oficial
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # 3. Datos del encabezado y resguardante
    hoy_str = datetime.date.today().strftime("%d/%m/%Y")
    ws["G5"] = hoy_str
    ws["A10"] = nombre_resguardante
    ws["C10"] = puesto_resguardante
    ws["E10"] = ""

    # 4. Ajustar filas si hay más de 14 activos
    num_activos = len(activos)
    template_slots = 14
    start_row = 14

    if num_activos > template_slots:
        extra_rows = num_activos - template_slots
        insert_at = start_row + template_slots
        ws.insert_rows(insert_at, amount=extra_rows)
        # Replicar estilos y altura de fila
        for r in range(insert_at, insert_at + extra_rows):
            ws.row_dimensions[r].height = ws.row_dimensions[start_row].height or 20
            for c in range(1, 8):
                src_cell = ws.cell(start_row, c)
                dest_cell = ws.cell(r, c)
                dest_cell.border = copy.copy(src_cell.border)
                dest_cell.font = copy.copy(src_cell.font)
                dest_cell.alignment = copy.copy(src_cell.alignment)

    # 5. Llenar los datos de los activos
    for idx, a in enumerate(activos):
        curr_row = start_row + idx
        num_inv = a.codigo_oficial if a.codigo_oficial else a.codigo_interno
        obs_val = a.observaciones or ("EN DESUSO" if a.estatus_activo == "EN_DESUSO" else "OPERATIVO")

        ws.cell(curr_row, 1, value=str(num_inv))
        ws.cell(curr_row, 2, value=str(a.descripcion or ""))
        ws.cell(curr_row, 3, value=str(a.especificacion or ""))
        ws.cell(curr_row, 4, value=str(a.numero_serie or "S/N"))
        ws.cell(curr_row, 5, value=str(a.marca or ""))
        ws.cell(curr_row, 6, value=str(a.modelo or ""))
        ws.cell(curr_row, 7, value=str(obs_val))

    # 6. Calcular celda de total de bienes
    total_row = start_row + max(num_activos, template_slots)
    ws.cell(total_row, 7, value=num_activos)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

