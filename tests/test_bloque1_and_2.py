import io
import openpyxl
from fastapi.testclient import TestClient
from backend.main import app
from backend.database import SessionLocal
from backend.models import Activo

def test_auto_code_generation_and_export():
    client = TestClient(app)
    created_ids = []

    # Obtener token de admin
    login_res = client.post("/api/auth/login", json={"username": "admin", "password": "Cobach3#Admin"})
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # Limpieza preventiva inicial de posibles remanentes
    db = SessionLocal()
    try:
        db.query(Activo).filter(
            Activo.descripcion.in_(["TEST ARTICULO GASTO AUTO", "TEST ARTICULO CA AUTO"])
        ).delete(synchronize_session=False)
        db.commit()
    finally:
        db.close()

    try:
        # 1. Test POST /api/activos sin codigo_interno para GASTO
        payload_gasto = {
            "descripcion": "TEST ARTICULO GASTO AUTO",
            "origen": "GASTO",
            "condicion_actual": "Buena 61% - 80%",
            "estatus_operativo": "OPERATIVO",
            "codigo_interno": None
        }
        res_gasto = client.post("/api/activos", json=payload_gasto, headers=headers)
        assert res_gasto.status_code == 201, f"Error: {res_gasto.text}"
        data_gasto = res_gasto.json()
        assert data_gasto["codigo_interno"].startswith("PL3-GTO-"), f"Unexpected code: {data_gasto['codigo_interno']}"
        created_ids.append(data_gasto["id"])

        # 2. Test POST /api/activos sin codigo_interno para CONTROL ADMINISTRATIVO
        payload_ca = {
            "descripcion": "TEST ARTICULO CA AUTO",
            "origen": "CONTROL ADMINISTRATIVO",
            "condicion_actual": "Buena 61% - 80%",
            "estatus_operativo": "OPERATIVO",
            "codigo_interno": None
        }
        res_ca = client.post("/api/activos", json=payload_ca, headers=headers)
        assert res_ca.status_code == 201, f"Error: {res_ca.text}"
        data_ca = res_ca.json()
        assert data_ca["codigo_interno"].startswith("PL3-CA-"), f"Unexpected code: {data_ca['codigo_interno']}"
        created_ids.append(data_ca["id"])

        # 3. Test Export Excel with custom columns: codigo_interno, descripcion, ubicacion
        res_exp = client.get("/api/export/excel?columnas=codigo_interno,descripcion,ubicacion", headers=headers)
        assert res_exp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res_exp.headers["content-type"]
        wb = openpyxl.load_workbook(io.BytesIO(res_exp.content))
        ws = wb.active
        header_row = [cell.value for cell in ws[1]]
        assert header_row == ["Código Interno", "Descripción", "Ubicación"], f"Headers mismatch: {header_row}"

        # 4. Test Export Excel default columns
        res_exp_all = client.get("/api/export/excel", headers=headers)
        assert res_exp_all.status_code == 200
        wb_all = openpyxl.load_workbook(io.BytesIO(res_exp_all.content))
        ws_all = wb_all.active
        header_row_all = [cell.value for cell in ws_all[1]]
        assert len(header_row_all) == 19, f"Expected 19 columns, got {len(header_row_all)}"

    finally:
        # Limpieza estricta e incondicional de los elementos de prueba
        if created_ids:
            clean_db = SessionLocal()
            try:
                clean_db.query(Activo).filter(Activo.id.in_(created_ids)).delete(synchronize_session=False)
                clean_db.commit()
            finally:
                clean_db.close()

    print("TEST BLOQUE 1 & 2 PASSED COMPLETELY WITHOUT LEFTOVERS!")

if __name__ == "__main__":
    test_auto_code_generation_and_export()
