import io
import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.main import app
from backend.database import SessionLocal
from backend.models import Activo, Usuario, BitacoraLog


@pytest.fixture
def auth_admin_client():
    client = TestClient(app)
    res = client.post("/api/auth/login", json={"username": "admin", "password": "Cobach3#Admin"})
    assert res.status_code == 200
    token = res.json()["access_token"]
    client.headers = {"Authorization": f"Bearer {token}"}
    return client


def test_soft_delete_and_restore(auth_admin_client):
    client = auth_admin_client

    # 1. Crear activo de prueba
    payload = {
        "descripcion": "ACTIVO PRUEBA SOFT DELETE",
        "origen": "GASTO",
        "condicion_actual": "Buena",
        "estatus_operativo": "OPERATIVO",
        "numero_serie": "TEST-SOFT-DEL-123"
    }
    create_res = client.post("/api/activos", json=payload)
    assert create_res.status_code == 201
    activo_id = create_res.json()["id"]
    cod_interno = create_res.json()["codigo_interno"]

    try:
        # 2. Verificar que aparece en lista normal
        list_res = client.get(f"/api/activos?q={cod_interno}")
        assert list_res.status_code == 200
        assert list_res.json()["total"] >= 1

        # 3. Soft Delete
        del_res = client.delete(f"/api/activos/{activo_id}")
        assert del_res.status_code == 200
        assert "movido a la papelera" in del_res.json()["message"]

        # 4. Verificar que YA NO aparece en lista normal
        list_after_del = client.get(f"/api/activos?q={cod_interno}")
        assert list_after_del.json()["total"] == 0

        # 5. Verificar que APARECE en papelera
        papelera_res = client.get("/api/activos/papelera/lista")
        assert papelera_res.status_code == 200
        p_items = papelera_res.json()["items"]
        matching = [p for p in p_items if p["id"] == activo_id]
        assert len(matching) == 1
        assert matching[0]["deleted_by"] == "admin"

        # 6. Verificar que se registró en Bitácora
        bitacora_res = client.get(f"/api/bitacora?q={cod_interno}")
        assert bitacora_res.status_code == 200
        b_items = bitacora_res.json()["items"]
        assert any(b["operacion"] == "ELIMINACION" for b in b_items)

        # 7. Restaurar activo
        restore_res = client.post(f"/api/activos/{activo_id}/restaurar")
        assert restore_res.status_code == 200
        assert "restaurado con éxito" in restore_res.json()["message"]

        # 8. Verificar que VUELVE a aparecer en lista normal y ya no en papelera
        list_restored = client.get(f"/api/activos?q={cod_interno}")
        assert list_restored.json()["total"] == 1

        papelera_after = client.get("/api/activos/papelera/lista")
        assert not any(p["id"] == activo_id for p in papelera_after.json()["items"])

    finally:
        # Purgar definitivamente
        client.delete(f"/api/activos/{activo_id}")
        client.delete(f"/api/activos/{activo_id}/permanente")


def test_generar_resguardo_oficial_excel(auth_admin_client):
    client = auth_admin_client

    # Tomar un lote de 2 activos existentes
    list_res = client.get("/api/activos?limit=2")
    assert list_res.status_code == 200
    items = list_res.json()["items"]
    ids = [i["id"] for i in items]

    res = client.post("/api/reportes/resguardo-oficial", json={"activo_ids": ids})
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    # Verificar plantilla institucional
    assert "COLEGIO DE BACHILLERES" in str(ws.cell(1, 7).value or "") or "COLEGIO DE BACHILLERES" in str(ws.cell(1, 1).value or "")
    assert ws.cell(13, 1).value == "Número de Inventario"
    assert ws.cell(13, 2).value == "Descripción"
    assert ws.cell(14, 1).value is not None


def test_bitacora_endpoint_filters(auth_admin_client):
    client = auth_admin_client
    res = client.get("/api/bitacora?limit=10")
    assert res.status_code == 200
    data = res.json()
    assert "total" in data
    assert "items" in data
    assert isinstance(data["items"], list)
    assert len(data["items"]) > 0
    # Verificar estructura del item
    first = data["items"][0]
    assert "operacion" in first
    assert "usuario_nombre" in first
    assert "fecha_hora" in first
